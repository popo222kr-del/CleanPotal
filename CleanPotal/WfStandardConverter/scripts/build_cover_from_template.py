"""정본 표지 + 재.개정 이력 시트 생성기 (공식 샘플 양식 복사 + 값 치환)."""
from copy import copy
import openpyxl
from openpyxl.styles import Border, Side

def _styles(ws, src_row, ncol):
    return {c: copy(ws.cell(src_row, c)._style) for c in range(2, ncol+1)}
def _apply(ws, r, styles, ncol):
    for c in range(2, ncol+1): ws.cell(r, c)._style = copy(styles[c])
def _merge(ws, r, segs):
    for a,b in segs: ws.merge_cells(f"{a}{r}:{b}{r}")
def _side(cell, top=None, bottom=None, left=None, right=None):
    b=cell.border
    cell.border=Border(top=Side(style=top) if top else b.top,
                       bottom=Side(style=bottom) if bottom else b.bottom,
                       left=Side(style=left) if left else b.left,
                       right=Side(style=right) if right else b.right)
def _box(ws, hr, last, c0, c1):
    if last<hr: return
    for c in range(c0,c1+1):
        _side(ws.cell(hr,c), top="medium"); _side(ws.cell(last,c), bottom="medium")
    for rr in range(hr,last+1):
        _side(ws.cell(rr,c0), left="medium"); _side(ws.cell(rr,c1), right="medium")
def _clear(ws, from_row, ncol=110):
    for m in list(ws.merged_cells.ranges):
        if m.min_row>=from_row: ws.unmerge_cells(str(m))
    for r in range(from_row, ws.max_row+1):
        for c in range(2, ncol+1):
            cell=ws.cell(r,c); cell.value=None; cell.style="Normal"

SUB_SEG=[("B","M"),("N","S"),("T","BW"),("BX","CF"),("CG","CY")]
HIS_SEG=[("B","M"),("N","S"),("T","AB"),("AC","CY")]
NCOLc=103
def populate_cover(ws, mno, mname, subs, history):
    st_title=copy(ws.cell(13,2)._style)
    st_subhdr=_styles(ws,14,NCOLc); st_subdat=_styles(ws,15,NCOLc)
    st_hishdr=_styles(ws,22,NCOLc); st_hisdat=_styles(ws,23,NCOLc)
    ws.cell(1,27,mname); ws.cell(1,86,mno); ws.cell(3,86,"Rev 00"); ws["CH2"]=None
    for a in ["X4","AG4","X5","X6","X8","AO5","BC5","BJ5","BQ5","BX5","CS5",
              "AO6","BC6","BJ6","BQ6","BX6","CS6","D10","D11","D12"]: ws[a]=None
    _clear(ws,13)
    for r in range(1, ws.max_row+1):
        for c in (105,106,107): ws.cell(r,c).value=None
    r=13
    ws.cell(r,2,"■ 표준 부속서 List"); ws.cell(r,2)._style=copy(st_title); ws.merge_cells(f"B{r}:Q{r}"); r+=1
    sub_hdr=r
    _apply(ws,r,st_subhdr,NCOLc); _merge(ws,r,SUB_SEG)
    ws.cell(r,2,"부속서 표준 No"); ws.cell(r,14,"Rev No"); ws.cell(r,20,"제목"); ws.cell(r,76,"담당자"); ws.cell(r,85,"비고"); ws.row_dimensions[r].height=17.2; r+=1
    sub_last=r-1
    for s in subs:
        _apply(ws,r,st_subdat,NCOLc); _merge(ws,r,SUB_SEG)
        ws.cell(r,2,s.get("no","")); ws.cell(r,14,s.get("rev","")); ws.cell(r,20,s.get("title",""))
        ws.cell(r,76,s.get("owner","")); ws.cell(r,85,s.get("note","")); ws.row_dimensions[r].height=17.2; sub_last=r; r+=1
    _box(ws, sub_hdr, sub_last, 2, NCOLc)
    r+=1
    his_title=r
    ws.cell(r,2,"■ 최종 표준 재.개정 내역"); ws.cell(r,2)._style=copy(st_title); ws.merge_cells(f"B{r}:P{r}"); r+=1
    his_hdr=r
    _apply(ws,r,st_hishdr,NCOLc); _merge(ws,r,HIS_SEG)
    ws.cell(r,2,"표준 No"); ws.cell(r,14,"Rev. No"); ws.cell(r,20,"일자"); ws.cell(r,29,"사유"); ws.row_dimensions[r].height=17.2; r+=1
    his_last=r-1
    for h in history:
        _apply(ws,r,st_hisdat,NCOLc); _merge(ws,r,HIS_SEG)
        ws.cell(r,2,h.get("stdno","")); ws.cell(r,14,h.get("rev","")); ws.cell(r,20,h.get("date","")); ws.cell(r,29,h.get("reason","")); ws.row_dimensions[r].height=17.2; his_last=r; r+=1
    _box(ws, his_hdr, his_last, 2, NCOLc)
    r+=1
    special_row=r
    ws.cell(r,2,"■ 특이사항"); ws.cell(r,2)._style=copy(st_title)
    # 사용자 지정 굵은 바깥쪽 박스
    _box(ws, 9, 12, 2, NCOLc)        # B9:CY12 표준 운영 목적
    # T12:W12 은 아래쪽만(좌/우/위 제거)
    _empty=Side(style=None)
    for c in range(20,24):
        b=ws.cell(12,c).border
        ws.cell(12,c).border=Border(bottom=b.bottom, top=_empty, left=_empty, right=_empty)
    _box(ws, his_title, his_title, 2, NCOLc)   # 최종재개정 제목행 (B24:CY24)
    _box(ws, special_row, special_row, 2, NCOLc)  # 특이사항 행 (B36:CY36)
    _box(ws, 1, special_row+6, 2, NCOLc)           # 전체 외곽 박스 (특이사항행+6까지, 패턴 상대참조)
    _e=Side(style=None)                            # B36:CY36 아래 테두리 제거
    for c in range(2, NCOLc+1):
        b=ws.cell(special_row,c).border
        ws.cell(special_row,c).border=Border(top=b.top, left=b.left, right=b.right, bottom=_e)

HH_SEG=[("B","C"),("E","G"),("I","O")]
HS_SEG=[("C","D"),("E","K"),("L","O")]
NCOLh=15
def populate_history(ws, blocks):
    st_title=copy(ws.cell(1,2)._style)
    st_blk=_styles(ws,2,NCOLh); st_sub=_styles(ws,3,NCOLh); st_rev=_styles(ws,4,NCOLh)
    title="■ "+(blocks[0].get("title") if blocks else "표준 재.개정 이력")
    _clear(ws,1,ncol=25)
    r=1
    ws.cell(r,2,title); ws.cell(r,2)._style=copy(st_title); ws.merge_cells(f"B{r}:O{r}"); ws.row_dimensions[r].height=30.8; r+=1
    for blk in blocks:
        hr=r
        _apply(ws,r,st_blk,NCOLh); _merge(ws,r,HH_SEG)
        ws.cell(r,2,blk.get("type","표준")); ws.cell(r,4,"No"); ws.cell(r,5,blk.get("stdno","")); ws.cell(r,8,"제목"); ws.cell(r,9,blk.get("title","")); ws.row_dimensions[r].height=17.2; r+=1
        _apply(ws,r,st_sub,NCOLh); _merge(ws,r,HS_SEG)
        ws.cell(r,2,"Rev. No"); ws.cell(r,3,"일자"); ws.cell(r,5,"사유"); ws.cell(r,12,"비고"); ws.row_dimensions[r].height=17.2; r+=1
        first_data=r; last_data=r-1
        for rv in blk.get("revs",[]):
            _apply(ws,r,st_rev,NCOLh); _merge(ws,r,HS_SEG)
            ws.cell(r,2,rv.get("rev","")); ws.cell(r,3,rv.get("date","")); ws.cell(r,5,rv.get("reason","")); ws.cell(r,12,rv.get("note","")); ws.row_dimensions[r].height=17.2; last_data=r; r+=1
        # 블록 외곽을 일관된 굵은(medium) 박스로 닫기: 좌(B)/우(O) 전체행, 마지막행 아래
        if last_data>=hr:
            for rr in range(hr, last_data+1):
                _side(ws.cell(rr,2), left="medium"); _side(ws.cell(rr,15), right="medium")
            for c in range(2,NCOLh+1): _side(ws.cell(last_data,c), bottom="medium")
        r+=1

def build_workbook(template_xlsx, cover_sheet, history_sheet, mno, mname, subs, history, blocks):
    wb=openpyxl.load_workbook(template_xlsx)
    for s in list(wb.sheetnames):
        if s not in (cover_sheet, history_sheet): del wb[s]
    try:
        for v in wb.views: v.firstSheet=0; v.activeTab=0
    except Exception: pass
    wb.active=0
    populate_cover(wb[cover_sheet], mno, mname, subs, history)
    populate_history(wb[history_sheet], blocks)
    wb[cover_sheet].title=("모표준_"+mno)[:31]
    return wb
# 구버전 호환
def build_cover(template_xlsx, template_sheet, mno, mname, subs, history):
    wb=openpyxl.load_workbook(template_xlsx)
    for s in list(wb.sheetnames):
        if s!=template_sheet: del wb[s]
    try:
        for v in wb.views: v.firstSheet=0; v.activeTab=0
    except Exception: pass
    wb.active=0
    populate_cover(wb[template_sheet], mno, mname, subs, history)
    return wb
